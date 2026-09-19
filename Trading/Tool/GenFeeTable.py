#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
费率表生成器（D-B · 费率生成式 SSOT · 2026-09-15）
==================================================

为什么有本文件
--------------
交接文档批判过一种反模式：**同一份数据存两份 + 配一套对账**（当时批的是
price_tick 与行情的那 122 行）。而落地费率时又重建了同一结构：xlsx 一份、
`Product.py` 手抄一份、再配快照对账 —— "手滑写错一个数"始终物理上可能。

本生成器把这条链路收成单向：

    券商费率表（Docs/手续费标准-.xlsx）
        → 本脚本解析
        → 重写 `Trading/Infra/Product.py` 内的 GENERATED 标记区块（纯数据，禁止手编）
        → Product.py 构造 Fee 时读同模块的 BASE / OVERRIDES，构造期注入档案

于是"手抄错一个数"在物理上不可能发生：`PRODUCT_PROFILES` 段里一个费率数字都没有。

**为什么是"标记区块"而不是独立文件**（用户拍板）
------------------------------------------------------------
最初实现为 `Trading/Infra/_FeeTable.py`（独立生成文件）。用户指出这与
交接文档冲突 —— 那一节把 `Trading/Infra/` 定死成 **7 个模块、按变异轴排**
（`Clock / Records / Period / Product / Instrument / StateDB / EventLog`），
且 §附C 原则 1 明令「名字 = 领域概念，不是技术机制」（`_` 前缀是可见性机制，
而该模块被跨模块引用，"私有"语义是假的）。

改为**区块内联**后：`Infra/` 仍是 7 个文件、零新增；机器生成内容与手写内容
靠**标记 + 逐字节护栏**分离，而不是靠文件边界。守卫分三层：
  ① 生成器只重写两块标记之间的文本，**区块外逐字节不动**；
  ② `--check` 重新解析 xlsx 与本文件区块逐字节比对，不一致 → rc=1；
  ③ `Trading/Test/test_product_fee_table.py [3f]/[3i]` 跑 ② 并**自检护栏会咬人**。

用法
----
    python Trading/Tool/GenFeeTable.py                    # 刷新 Product.py 的费率区块
    python Trading/Tool/GenFeeTable.py --check            # 只校验，不写盘（CI / 测试用）
    python Trading/Tool/GenFeeTable.py --target <path>    # 指定目标文件（测试自检用）

`--check` 语义：重新解析 xlsx，与目标文件区块逐字节比对（忽略 CRLF），不一致 → rc=1
并打印差异摘要（这是"xlsx 改了但没重新生成"以及"有人手改了区块"的双重护栏）。
xlsx 不存在时退出码 0 并打印 skip（离线/裁剪仓库不该因此变红）。

解析口径（与 xlsx 排版绑定，改排版必同步改这里）
------------------------------------------------
xlsx `Sheet1` 的 B 列是品种中文名、D 列是费率文本。一个品种块 = 从 B 列非空那行
开始、到下一个 B 列非空行为止的所有非空 D 值：

  · 块内**第一个非"交割"值 = 开仓档**（= 平昨档，xlsx 从不单列平昨）；
  · 块内**第二个非"交割"值 = 平今档**（xlsx 常省略"平今"字样，靠位置表达，
    如铜块 D150="0.5%%" / D151="1%%"）；
  · 第三个及以后（"交割X"）= 不参与交割，忽略。

费率文本 → (kind, value)：
  · "0.23%%" / "交易0.23%%"   → ("rate", 0.23)     万分之几
  · "平今2.3%%"               → ("rate", 2.3)
  · "平今免" / "平今免收"      → ("per_lot", 0.0)    免收
  · "10" / "平今60"            → ("per_lot", 10.0)   元/手
  · "交割0.5%%"                → 忽略

覆盖档（合约月份差异化费率）：B 列形如「黄金6、12合约&2607、2608、2609、2610合约」，
即「品种名 + 含"合约"的月份描述」→ 记为该品种的覆盖档，**当前不消费**（字段位保留，
见 Product.fee_overrides）。期权行（「黄金期权」）因不含"合约"二字被排除。
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from typing import Dict, List, Optional, Tuple

_HERE = os.path.dirname(os.path.abspath(__file__))          # <root>/Trading/Tool
_ROOT = os.path.dirname(os.path.dirname(_HERE))             # <root>

XLSX_REL = "Docs/手续费标准-.xlsx"
TARGET_REL = "Trading/Infra/Product.py"
SHEET = "Sheet1"

# 标记行（**生成器与测试共用的契约**；改这两行 = 改契约，两处一起改）
BLOCK_BEGIN = "# >>> GENERATED: FeeTable（禁止手工编辑；改费率请改 xlsx 后重跑生成器）"
BLOCK_END = "# <<< END GENERATED"

# 品种中文名 → 档案键（新增品种必须同步这里 + Product.PRODUCT_PROFILES）
NAME_TO_CODE: Dict[str, str] = {
    "沪深300指数": "IF",
    "上证50": "IH",
    "中证500": "IC",
    "中证1000": "IM",
    "黄金": "AU",
    "白银": "AG",
    "铜": "CU",
    "PTA": "TA",
}

# 费率文本前缀（长的在前，避免 "交易" 吃掉 "交易日…" 之类；本表用例无此冲突）
_PREFIXES = ("隔日开平", "当日开平", "交易", "平今", "交割")


def _parse_fee(text) -> Optional[Tuple[Optional[str], Tuple[str, float]]]:
    """费率文本 → (tag, (kind, value))。tag ∈ {None, "交易", "平今", "交割", …}。"""
    s = str(text or "").strip()
    if not s:
        return None
    tag: Optional[str] = None
    for pre in _PREFIXES:
        if s.startswith(pre):
            tag = pre
            s = s[len(pre):].strip()
            break
    if "免" in s:                       # "平今免" / "平今免收"
        return tag, ("per_lot", 0.0)
    if s.endswith("%%"):
        return tag, ("rate", float(s[:-2]))
    if s.endswith("%"):
        return tag, ("rate", float(s[:-1]))
    return tag, ("per_lot", float(s))


def _classify(name: str) -> Tuple[Optional[str], bool]:
    """B 列文本 → (品种代码, 是否覆盖档)。不认识的品种 → (None, False)。"""
    n = str(name or "").strip()
    for cn, code in NAME_TO_CODE.items():
        if n == cn:
            return code, False
        # 覆盖档 = 「品种名 + 含"合约"的月份描述」；期权行不含"合约" → 排除
        if n.startswith(cn) and ("合约" in n):
            return code, True
    return None, False


def build_tables(xlsx_path: str) -> Dict[str, object]:
    """解析 xlsx → {"base": {code: (开仓, 平今|None)}, "overrides": {code: [(标签, 开仓, 平今|None), …]}}。

    费率统一表达为 (kind, value) 元组（rate=万分之几 / per_lot=元每手），
    与 `Infra/Product.py` 的 `Fee` 字段一一对应；本函数**不 import Fee**，
    避免生成器与生产代码互相绑定。
    """
    import openpyxl  # 延迟导入：只有真正需要解析时才要求该依赖

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError("xlsx 缺少工作表 {!r}（现有: {}）".format(
            SHEET, ", ".join(wb.sheetnames)))
    ws = wb[SHEET]

    blocks: List[Dict[str, object]] = []
    cur: Optional[Dict[str, object]] = None
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        d = ws.cell(r, 4).value
        if b is not None and str(b).strip():
            code, is_ov = _classify(str(b).strip())
            # 认得 → 开新块；认不得（含期权行 / 段落里的杂项）→ 关闭当前块
            cur = ({"code": code, "override": is_ov,
                    "name": str(b).strip(), "rows": []} if code else None)
            if cur is not None:
                blocks.append(cur)
        if cur is not None and d is not None and str(d).strip():
            parsed = _parse_fee(d)
            if parsed is not None:
                cur["rows"].append((r, parsed[0], parsed[1]))

    base: Dict[str, Tuple[Tuple[str, float], Optional[Tuple[str, float]]]] = {}
    overrides: Dict[str, List[Tuple[str, Tuple[str, float], Optional[Tuple[str, float]]]]] = {}
    labels: Dict[str, str] = {}

    for blk in blocks:
        code = blk["code"]
        rows = blk["rows"]
        # 交割档不参与（本系统不交割）；其余按位置：第 1 = 开仓、第 2 = 平今
        usable = [(rn, fee) for rn, tag, fee in rows if tag != "交割"]
        if not usable:
            continue
        o = usable[0][1]
        ct = usable[1][1] if len(usable) > 1 else None
        if blk["override"]:
            label = str(blk["name"])
            for cn in NAME_TO_CODE:
                if label.startswith(cn):
                    label = label[len(cn):]
                    break
            overrides.setdefault(code, []).append((label, o, ct))
        else:
            base[code] = (o, ct)
            labels[code] = str(blk["name"])
    return {"base": base, "overrides": overrides, "labels": labels}


# ══════════════════════════════════════════════════════════════════
# 渲染（区块内容 —— 纯数据，无时间戳，保证幂等可复现）
# ══════════════════════════════════════════════════════════════════
def _fmt_fee(fee) -> str:
    if fee is None:
        return "None"
    return "({!r}, {!r})".format(fee[0], fee[1])


def render_block(tables: Dict[str, object], nl: str = "\n") -> str:
    """渲染**两块标记之间**的正文（不含标记行本身，不加尾部换行）。

    幂等：同一份 tables 两次调用逐字节相同（无时间戳、无随机顺序 —— 全按代码排序）。
    """
    base = tables["base"]
    overrides = tables["overrides"]
    labels = tables["labels"]
    L: List[str] = []
    L.append("# 纯数据元组（kind, value）：")
    L.append('#   kind="rate"    → value = 万分之几（xlsx 里打印的 "0.23%%" 就存 0.23）')
    L.append('#   kind="per_lot" → value = 元/手')
    L.append("FeeTuple = Tuple[str, float]")
    L.append("Pair = Tuple[FeeTuple, Optional[FeeTuple]]   # (开仓档, 平今档; None=同开仓档)")
    L.append("")
    L.append("# 源表品种名（备案：便于人工回查 xlsx 行）")
    L.append("BASE_LABELS: Dict[str, str] = {")
    for code in sorted(labels):
        L.append("    {!r}: {!r},".format(code, labels[code]))
    L.append("}")
    L.append("")
    L.append("# 基准档：品种键 → (开仓/平昨档, 平今档)。平今 None = xlsx 无独立平今行（同开仓档）。")
    L.append("BASE: Dict[str, Pair] = {")
    for code in sorted(base):
        o, ct = base[code]
        L.append("    {!r}: ({}, {}),".format(code, _fmt_fee(o), _fmt_fee(ct)))
    L.append("}")
    L.append("")
    L.append("# 覆盖档（合约月份差异化费率）：品种键 → ((标签, 开仓档, 平今档), …)")
    L.append("# ⚠️ 当前**不消费**（Product.fee_overrides 字段位保留、默认空）。")
    L.append("#    AU/AG 主力滚到覆盖档合约时若不启用本表，回测成本会静默低估；")
    L.append("#    启用方式见 Product.fee_overrides 字段注释（一处收口，勿在别处另开分支）。")
    L.append("OVERRIDES: Dict[str, List[OverrideItem]] = {")
    for code in sorted(overrides):
        items = overrides[code]
        parts = ", ".join("({!r}, {}, {})".format(lb, _fmt_fee(o), _fmt_fee(ct))
                          for lb, o, ct in items)
        L.append("    {!r}: [{}],".format(code, parts))
    L.append("}")
    return nl.join(L)


# 标记行匹配（允许行首尾空白；必须整行命中）
# ⚠️ 收尾用 (?=\r?$) 而非 $：_read() 以 newline="" 读盘、原样保留换行，
#    CRLF 文件里 "$" 会落在 "\r" 之后 → 整行失配、区块报"找不到"。前瞻不吞
#    "\r"（m1.end()/m2.start() 偏移不变，替换后文件仍是 CRLF），LF 文件照旧。
_BEGIN_RE = re.compile(r"^[ \t]*" + re.escape(BLOCK_BEGIN) + r"[ \t]*(?=\r?$)", re.M)
_END_RE = re.compile(r"^[ \t]*" + re.escape(BLOCK_END) + r"[ \t]*(?=\r?$)", re.M)


def detect_nl(text: str) -> str:
    """目标文件用的换行符（生成器必须沿用，不能把 CRLF 文件改成 LF）。"""
    return "\r\n" if "\r\n" in text else "\n"


def replace_block(text: str, body: str, nl: str) -> str:
    """把两块标记之间的正文替换成 body。**区块外逐字节不动**（含标记行本身）。

    标记缺失或出现多次 → ValueError（宁可直接报错，也不要"以为替换成功"）。
    """
    m1 = _BEGIN_RE.search(text)
    m2 = _END_RE.search(text)
    if not m1 or not m2 or m2.start() <= m1.start():
        raise ValueError("目标文件缺少费率区块标记（或顺序颠倒）：需要 {!r} … {!r}".format(
            BLOCK_BEGIN, BLOCK_END))
    if _BEGIN_RE.search(text, m1.end()) or _END_RE.search(text, m2.end()):
        raise ValueError("费率区块标记出现多次 —— 请人工确认后再重跑")
    # head 含 BEGIN 标记行；tail 从 END 标记行开始
    head = text[:m1.end()]
    tail = text[m2.start():]
    return head + nl + body + nl + tail


def _read(path: str) -> str:
    with open(path, encoding="utf-8", newline="") as f:
        return f.read()


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="从券商费率表刷新 {} 内的 GENERATED 费率区块".format(TARGET_REL))
    ap.add_argument("--check", action="store_true",
                    help="只校验目标文件区块与 xlsx 是否一致，不写盘")
    ap.add_argument("--xlsx", default=os.path.join(_ROOT, XLSX_REL),
                    help="费率表路径（默认仓库内 {!r}）".format(XLSX_REL))
    ap.add_argument("--target", default=None,
                    help="目标文件（默认仓库内 {!r}；测试自检用）".format(TARGET_REL))
    args = ap.parse_args(argv)

    target = args.target or os.path.join(_ROOT, TARGET_REL)
    if not os.path.isfile(target):
        print("[GenFeeTable] ✗ 找不到目标文件 {} —— 无法校验/写入".format(target))
        return 2
    if not os.path.isfile(args.xlsx):
        msg = "[GenFeeTable] 找不到费率表 {} —— 跳过（离线/裁剪仓库不算错误）".format(args.xlsx)
        print(msg)
        return 0

    try:
        tables = build_tables(args.xlsx)
    except ImportError:
        print("[GenFeeTable] 缺少 openpyxl，无法解析 xlsx —— 跳过")
        return 0 if args.check else 2

    missing = sorted(set(NAME_TO_CODE.values()) - set(tables["base"]))
    if missing:
        print("[GenFeeTable] ✗ 源表里没有这些品种的基准档: {}".format(", ".join(missing)))
        return 2

    text = _read(target)
    nl = detect_nl(text)
    body = render_block(tables, nl)

    try:
        new_text = replace_block(text, body, nl)
    except ValueError as e:
        print("[GenFeeTable] ✗ {}".format(e))
        return 1

    if args.check:
        if new_text == text:
            print("[GenFeeTable] ✓ 一致：{} 的费率区块与 {} 同步".format(
                os.path.relpath(target, _ROOT).replace("\\", "/"), XLSX_REL))
            return 0
        print("[GenFeeTable] ✗ {} 的费率区块与 {} **不一致** —— "
              "请重跑生成器（或区块被手改过，请还原）".format(
                  os.path.relpath(target, _ROOT).replace("\\", "/"), XLSX_REL))
        return 1

    if new_text != text:
        with open(target, "w", encoding="utf-8", newline="") as f:
            f.write(new_text)
        print("[GenFeeTable] ✓ 已刷新 {} 的费率区块（{} 个基准档，{} 个带覆盖档）".format(
            os.path.relpath(target, _ROOT).replace("\\", "/"),
            len(tables["base"]), len(tables["overrides"])))
        for code in sorted(tables["base"]):
            o, ct = tables["base"][code]
            print("    {:<3s} 开仓 {} | 平今 {}".format(code, _fmt_fee(o), _fmt_fee(ct)))
    else:
        print("[GenFeeTable] ✓ {} 的费率区块已是最新（无改动）".format(
            os.path.relpath(target, _ROOT).replace("\\", "/")))
    return 0


if __name__ == "__main__":
    sys.exit(main())
